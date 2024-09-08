import json

from docx import Document
import pandas
from docx2python import docx2python
from selenium.common.exceptions import NoSuchElementException
class SeCommands():
    def __init__(self,driver):
        self.driver = driver

    def open_site(self, login, password):
        self.driver.get('https://nz.ua/')
        self.driver.click('button:contains("Увійти")')
        self.driver.type('#loginform-login', login)
        self.driver.sleep(0.05)
        self.driver.type('#loginform-password', password)
        self.driver.click('button.form-submit-btn')
        self.driver.get('https://nz.ua/journal/list')
        try:
            self.driver.find_element(by='css selector', value='.header-customer__name')
            return True
        except NoSuchElementException:
            return False


    def get_subject(self):
        elements = self.driver.find_elements('xpath', "//table[@class='journal-choose']//tr//td[1]")
        return elements
    def get_classes(self,subject):
        selected_subject = subject
        class_row = self.driver.find_elements('xpath',
                                              f"//td[normalize-space()='{selected_subject}']/following-sibling::td/a")
        class_list = [element.text.strip() for element in class_row if element.text.strip()]
        return class_list

    def connect_journal(self, combo_box1, combo_box2):
        selected_subject = combo_box1.currentText()
        selected_class = combo_box2.currentText()

        class_link = self.driver.find_element('xpath',
                                              f"//td[normalize-space()='{selected_subject}']/following-sibling::td/a[contains(text(), '{selected_class}')]")
        class_link.click()
    def journal_info(self):

        element=self.driver.find_element(by='css selector', value='.journal-scores__title')
        return element.text
    def is_cell_merged(self,cell):
        tc = cell._tc
        grid_span = tc.xpath('.//w:gridSpan')
        v_merge = tc.xpath('.//w:vMerge')

        if grid_span or v_merge:
            return True

        return False
    

    def find_table_docx(self,document_path):
        num=0
        max_lenght=0
        docm = docx2python(document_path)
        for id,table in enumerate(docm.body):
            if len(table) > max_lenght:
                num=id
                return num
            

    def find_row_docx(self,document_path,start_number):
        num=0
        max_lenght=0
        docm = docx2python(document_path)
        for id,table in enumerate(docm.body):
            if len(table) > max_lenght:
                max_lenght=len(table)
                num=id
        first_table = docm.body[num]    
        for idx,row in enumerate(first_table):
            number=row[0][0].strip()
            if '.' in number or ')' in number:
                number = number[:-1]
            if number == f'{start_number}':
                return idx



    def add_data(self,document_path,engine,lesson:int,column_number=None,homework=False):

        if engine == ['docx','xlrd']:
            lesson_number = int(self.find_row_docx(document_path=document_path,start_number=lesson))
            doc = Document(document_path)
            table = doc.tables[0]
            x=0
            excel_df = None
        else:
            excel_df = pandas.read_excel(document_path, engine=engine)
            number = excel_df[excel_df['№'] == lesson].index[0]

        ele=self.driver.find_element(by='tag name',value='thead').find_elements(by='class name',value='pt-point')
        chans = ele
        for i in range(len(chans)):
            homework_table = self.driver.find_elements(by='css selector', value='.homework-row')[1:]
            chan=self.driver.find_elements(by='css selector', value='div.homework__item a.modal-box')
            self.driver.execute_script("arguments[0].scrollIntoView(true);", homework_table[-1])
            self.driver.execute_script("arguments[0].scrollIntoView(true);", chan[i])
            homework_items = homework_table[i].find_elements(by='css selector', value='.homework__item')[1]
            self.driver.sleep(0.5)
            if not homework_items.text.strip():
                element = chan[i]
                element.click()
                self.driver.sleep(0.5)

                if engine == 'docx':
                    if self.is_cell_merged(table.cell(lesson_number + i+ x, column_number - 1)):
                        x+=1
                    cell_text = table.cell(lesson_number + i+ x, column_number - 1).text
                else:
                    cell_text=excel_df.iloc[number+i,2]
                self.driver.type('#osvitaschedulereal-lesson_topic',cell_text)
                self.driver.type('#osvitaschedulereal-lesson_number_in_plan', str(lesson + i))
                if i != len(chans)-1:
                    self.driver.click("//select[@id='osvitaschedulereal-hometask_to']/option[2]")
                if not engine == 'docx' and homework:
                    self.driver.type('#osvitaschedulereal-hometask',excel_df.iloc[number+i].iloc[3])
                self.driver.sleep(1)
                self.driver.execute_script("window.scrollTo(0, document.body.scrollHeight);")
                self.driver.sleep(1)
                button = self.driver.find_element(by='css selector',value='.submit .form-submit-btn')
                self.driver.sleep(0.1)
                button.click()
                self.driver.sleep(1)
            else:
                print(f"Урок {i + 1} заповнений!")
        print("Завершено")

    def back_to_journals(self):
        self.driver.get('https://nz.ua/journal/list')
    def close(self):
        if self.driver:
            self.driver.quit()
